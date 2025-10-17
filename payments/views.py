from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponseForbidden, HttpResponse
from .models import Course, Student, Payment, CustomUser
from .forms import CourseForm, StudentForm, PaymentForm
from django.db.models import Q, Sum, Count
from django.db.models.functions import TruncMonth
import json
import os
from django.conf import settings
from datetime import datetime
from docx import Document
import openpyxl
from .forms import CustomUserEditForm,CustomUserCreationForm

# ====== Главная страница ======
@login_required
def index(request):
    return render(request, "payments/dashboard.html")


# ====== Курсы ======
@login_required
def course_list(request):
    user = request.user
    if user.role == 'student':
        payments = Payment.objects.filter(student__user=user)
        courses = [p.course for p in payments]
    else:
        courses = Course.objects.all()
    return render(request, "payments/courses_list.html", {"courses": courses})

@login_required
def course_create(request):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для добавления курсов.")
    if request.method == "POST":
        form = CourseForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect("course_list")
    else:
        form = CourseForm()
    return render(request, "payments/course_form.html", {"form": form, "title": "Добавить курс"})

@login_required
def course_edit(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для редактирования курсов.")
    course = get_object_or_404(Course, pk=pk)
    if request.method == "POST":
        form = CourseForm(request.POST, instance=course)
        if form.is_valid():
            form.save()
            return redirect("course_list")
    else:
        form = CourseForm(instance=course)
    return render(request, "payments/course_form.html", {"form": form, "title": "Редактировать курс"})

@login_required
def course_delete(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для удаления курсов.")
    course = get_object_or_404(Course, pk=pk)
    if request.method == "POST":
        course.delete()
        return redirect("course_list")
    return render(request, "payments/course_confirm_delete.html", {"course": course})


# ====== Пользователи ======
@login_required
def user_list(request):
    if request.user.role != 'admin':
        return HttpResponseForbidden("Только админ может управлять пользователями.")
    users = CustomUser.objects.all()
    return render(request, "payments/user_list.html", {"users": users})

@login_required
def user_create(request):
    if request.user.role != 'admin':
        return HttpResponseForbidden("Только админ может создавать пользователей.")
    
    if request.method == "POST":
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect("user_list")
    else:
        form = CustomUserCreationForm()
    
    return render(request, "payments/user_create.html", {"form": form})

@login_required
def user_edit(request, pk):
    if request.user.role != 'admin':
        return HttpResponseForbidden("Только админ может редактировать пользователей.")
    
    user = get_object_or_404(CustomUser, pk=pk)
    
    if request.method == "POST":
        form = CustomUserEditForm(request.POST, instance=user)
        if form.is_valid():
            form.save()
            return redirect("user_list")
    else:
        form = CustomUserEditForm(instance=user)
    
    return render(request, "payments/user_edit.html", {"form": form})



# ====== Студенты ======
@login_required
def student_list(request):
    if request.user.role == 'student':
        # Студент видит только себя
        users = CustomUser.objects.filter(id=request.user.id)
    else:
        # Менеджер и админ видят всех студентов
        users = CustomUser.objects.filter(role="student")
    students = Student.objects.filter(user__in=users)
    return render(request, "payments/student_list.html", {"students": students})

@login_required
def student_create(request):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для добавления студента.")
    if request.method == "POST":
        form = StudentForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect("student_list")
    else:
        form = StudentForm()
    return render(request, "payments/student_form.html", {"form": form, "title": "Добавить студента"})

@login_required
def student_edit(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для редактирования студента.")
    student = get_object_or_404(Student, pk=pk)
    form = StudentForm(request.POST or None, instance=student)
    if form.is_valid():
        form.save()
        return redirect("student_list")
    return render(request, "payments/student_form.html", {"form": form, "title": "Редактировать студента"})

@login_required
def student_delete(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для удаления студента.")
    student = get_object_or_404(Student, pk=pk)
    if request.method == "POST":
        student.delete()
        return redirect("student_list")
    return render(request, "payments/student_confirm_delete.html", {"student": student})


# ====== Оплаты ======
def payment_list(request):
    user = request.user
    if hasattr(user, 'role') and user.role == 'student':
        student = Student.objects.filter(user=user).first()
        if not student:
            # Студент без профиля — показываем пустой список или сообщение
            payments = Payment.objects.none()
        else:
            payments = Payment.objects.filter(student=student)
    else:
        payments = Payment.objects.all()
    
    query = request.GET.get('q', '')
    course_id = request.GET.get('course', '')
    status = request.GET.get('status', '')

    if query:
        payments = payments.filter(
            Q(student__user__first_name__icontains=query) |
            Q(student__user__last_name__icontains=query)
        )
    if course_id:
        payments = payments.filter(course_id=course_id)
    if status:
        payments = payments.filter(status=status)

    courses = Course.objects.all()

    return render(request, 'payments/payment_list.html', {
        'payments': payments,
        'courses': courses,
        'query': query,
        'selected_course': course_id,
        'selected_status': status,
    })


@login_required
def payment_create(request):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для создания оплаты.")
    form = PaymentForm(request.POST or None)
    if form.is_valid():
        form.save()
        return redirect('payment_list')
    return render(request, 'payments/payment_form.html', {'form': form, 'title': 'Добавить оплату'})

@login_required
def payment_edit(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для редактирования оплаты.")
    payment = get_object_or_404(Payment, pk=pk)
    form = PaymentForm(request.POST or None, instance=payment)
    if form.is_valid():
        form.save()
        return redirect('payment_list')
    return render(request, 'payments/payment_form.html', {'form': form, 'title': 'Редактировать оплату'})

@login_required
def payment_delete(request, pk):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для удаления оплаты.")
    payment = get_object_or_404(Payment, pk=pk)
    if request.method == "POST":
        payment.delete()
        return redirect('payment_list')
    return render(request, 'payments/payment_confirm_delete.html', {'payment': payment})

@login_required
def generate_contract(request, payment_id):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для генерации договора.")
    payment = get_object_or_404(Payment, pk=payment_id)
    template_path = os.path.join(settings.BASE_DIR, 'static', 'contracts', 'contract_template.docx')
    doc = Document(template_path)
    for p in doc.paragraphs:
        p.text = p.text.replace('{{ contract_number }}', str(payment.id))
        p.text = p.text.replace('{{ date }}', datetime.now().strftime('%d.%m.%Y'))
        p.text = p.text.replace('{{ student_name }}', payment.student.user.get_full_name())
        p.text = p.text.replace('{{ course_name }}', payment.course.name)
        p.text = p.text.replace('{{ start_date }}', payment.course.start_date.strftime('%d.%m.%Y'))
        p.text = p.text.replace('{{ end_date }}', payment.course.end_date.strftime('%d.%m.%Y'))
        p.text = p.text.replace('{{ price }}', str(payment.amount))
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    filename = f"contract_{payment.student.user.username}_{payment.course.name}.docx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    doc.save(response)
    return response

@login_required
def export_payments_excel(request):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для экспорта данных.")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Оплаты"
    ws.append(['Студент', 'Курс', 'Сумма', 'Дата'])
    for payment in Payment.objects.select_related('student__user', 'course'):
        ws.append([
            payment.student.user.get_full_name(),
            payment.course.name,
            payment.amount,
            payment.date.strftime('%d.%m.%Y')
        ])
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=payments.xlsx'
    wb.save(response)
    return response

@login_required
def payments_stats(request):
    if request.user.role not in ['manager', 'admin']:
        return HttpResponseForbidden("Нет прав для просмотра статистики.")
    course_payments = Payment.objects.values('course__name').annotate(total=Sum('amount'))
    course_labels = [item['course__name'] for item in course_payments]
    course_values = [float(item['total']) for item in course_payments]
    monthly = Payment.objects.annotate(month=TruncMonth('date')).values('month').annotate(count=Count('id')).order_by('month')
    month_labels = [item['month'].strftime('%Y-%m') for item in monthly]
    month_values = [item['count'] for item in monthly]
    return render(request, 'payments/payments_stats.html', {
        'course_labels': json.dumps(course_labels),
        'course_values': json.dumps(course_values),
        'month_labels': json.dumps(month_labels),
        'month_values': json.dumps(month_values),
    })
